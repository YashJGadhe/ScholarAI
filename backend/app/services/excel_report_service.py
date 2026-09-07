"""Excel report generation using openpyxl.

Creates institutional-format Excel reports with:
- Grouped platform headers (merged cells)
- Year/month columns under each platform
- Total and Average rows
- Styled headers, borders, freeze panes
- Proper column widths
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PLATFORM_COLORS = {
    "WOS": "1299B8",
    "SCOPUS": "E9711C",
    "GOOGLE_SCHOLAR": "3D6DE0",
}

PLATFORM_LABELS = {
    "WOS": "Web of Science",
    "SCOPUS": "Scopus",
    "GOOGLE_SCHOLAR": "Google Scholar",
}

MONTH_NAMES = [
    "", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
    "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
]


def _style_header(cell, fill_color: str, bold: bool = True):
    cell.font = Font(bold=bold, size=11, color="FFFFFF")
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )


def _style_cell(cell, bold: bool = False, align: str = "center"):
    cell.font = Font(bold=bold, size=10)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )


def generate_year_wise_excel(report: dict) -> bytes:
    """Generate Excel file for year-wise report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Year-Wise Citations"

    platforms = report["platforms"]
    years = report["years"]
    rows = report["rows"]
    totals = report["totals"]
    averages = report["averages"]
    dept = report["department"]

    # Title rows
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(platforms) * (len(years) + 1))
    title_cell = ws.cell(row=1, column=1, value=f"{dept} Department")
    title_cell.font = Font(bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2 + len(platforms) * (len(years) + 1))
    subtitle = ws.cell(row=2, column=1, value=f"Citation Report — {report['start_year']} to {report['end_year']}")
    subtitle.font = Font(bold=True, size=12, color="555555")
    subtitle.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=2 + len(platforms) * (len(years) + 1))
    gen = ws.cell(row=3, column=1, value=f"Generated: {datetime.now().strftime('%d %B %Y')}")
    gen.font = Font(size=10, italic=True, color="777777")
    gen.alignment = Alignment(horizontal="center")

    # Header row 1: Platform groups (row 5)
    header_row = 5
    ws.cell(row=header_row, column=1, value="Sr. No.")
    _style_header(ws.cell(row=header_row, column=1), "29405C")
    ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row + 1, end_column=1)

    ws.cell(row=header_row, column=2, value="Faculty Name")
    _style_header(ws.cell(row=header_row, column=2), "29405C")
    ws.merge_cells(start_row=header_row, start_column=2, end_row=header_row + 1, end_column=2)

    col_offset = 3
    for plat in platforms:
        color = PLATFORM_COLORS.get(plat, "29405C")
        label = PLATFORM_LABELS.get(plat, plat)
        num_cols = len(years) + 1  # years + total
        ws.merge_cells(
            start_row=header_row,
            start_column=col_offset,
            end_row=header_row,
            end_column=col_offset + num_cols - 1,
        )
        cell = ws.cell(row=header_row, column=col_offset, value=f"Citations in {label}")
        _style_header(cell, color)
        col_offset += num_cols

    # Header row 2: Year columns (row 6)
    sub_header_row = 6
    col_offset = 3
    for plat in platforms:
        color = PLATFORM_COLORS.get(plat, "29405C")
        for year in years:
            cell = ws.cell(row=sub_header_row, column=col_offset, value=year)
            _style_header(cell, color)
            col_offset += 1
        cell = ws.cell(row=sub_header_row, column=col_offset, value="Total")
        _style_header(cell, color, bold=True)
        col_offset += 1

    # Data rows
    data_start_row = 7
    for i, row in enumerate(rows):
        r = data_start_row + i
        ws.cell(row=r, column=1, value=row["sr_no"])
        _style_cell(ws.cell(row=r, column=1))

        ws.cell(row=r, column=2, value=row["faculty_name"])
        _style_cell(ws.cell(row=r, column=2), align="left")

        col_offset = 3
        for plat in platforms:
            plat_data = row.get(plat, {})
            for year in years:
                val = plat_data.get(year)
                cell = ws.cell(row=r, column=col_offset, value="NA" if val is None else val)
                _style_cell(ws.cell(row=r, column=col_offset))
                col_offset += 1
            total_val = plat_data.get("total", 0)
            cell = ws.cell(row=r, column=col_offset, value=total_val)
            _style_cell(ws.cell(row=r, column=col_offset), bold=True)
            col_offset += 1

    # Total row
    total_row = data_start_row + len(rows)
    ws.cell(row=total_row, column=1, value="")
    _style_cell(ws.cell(row=total_row, column=1))
    ws.cell(row=total_row, column=2, value="Total")
    _style_cell(ws.cell(row=total_row, column=2), bold=True, align="left")

    col_offset = 3
    for plat in platforms:
        plat_totals = totals.get(plat, {})
        for year in years:
            val = plat_totals.get(year, 0)
            cell = ws.cell(row=total_row, column=col_offset, value=val)
            _style_cell(ws.cell(row=total_row, column=col_offset), bold=True)
            col_offset += 1
        cell = ws.cell(row=total_row, column=col_offset, value=plat_totals.get("total", 0))
        _style_cell(ws.cell(row=total_row, column=col_offset), bold=True)
        col_offset += 1

    # Average row
    avg_row = total_row + 1
    ws.cell(row=avg_row, column=1, value="")
    _style_cell(ws.cell(row=avg_row, column=1))
    ws.cell(row=avg_row, column=2, value="Average")
    _style_cell(ws.cell(row=avg_row, column=2), bold=True, align="left")

    col_offset = 3
    for plat in platforms:
        plat_avgs = averages.get(plat, {})
        for year in years:
            val = plat_avgs.get(year, 0)
            cell = ws.cell(row=avg_row, column=col_offset, value=val)
            _style_cell(ws.cell(row=avg_row, column=col_offset), bold=True)
            col_offset += 1
        cell = ws.cell(row=avg_row, column=col_offset, value=plat_avgs.get("total", 0))
        _style_cell(ws.cell(row=avg_row, column=col_offset), bold=True)
        col_offset += 1

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    for c in range(3, col_offset):
        ws.column_dimensions[get_column_letter(c)].width = 10

    # Freeze panes
    ws.freeze_panes = "C7"

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_month_wise_excel(report: dict) -> bytes:
    """Generate Excel file for month-wise report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Month-Wise Citations"

    platforms = report["platforms"]
    months = report["months"]
    rows = report["rows"]
    totals = report["totals"]
    averages = report["averages"]
    dept = report["department"]
    year = report["year"]

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(platforms) * (len(months) + 1))
    title_cell = ws.cell(row=1, column=1, value=f"{dept} Department")
    title_cell.font = Font(bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center")

    start_m = MONTH_NAMES[report["start_month"]]
    end_m = MONTH_NAMES[report["end_month"]]
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2 + len(platforms) * (len(months) + 1))
    subtitle = ws.cell(row=2, column=1, value=f"Citation Report — {start_m} to {end_m} {year}")
    subtitle.font = Font(bold=True, size=12, color="555555")
    subtitle.alignment = Alignment(horizontal="center")

    # Headers
    header_row = 4
    ws.cell(row=header_row, column=1, value="Sr. No.")
    _style_header(ws.cell(row=header_row, column=1), "29405C")
    ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row + 1, end_column=1)

    ws.cell(row=header_row, column=2, value="Faculty Name")
    _style_header(ws.cell(row=header_row, column=2), "29405C")
    ws.merge_cells(start_row=header_row, start_column=2, end_row=header_row + 1, end_column=2)

    col_offset = 3
    for plat in platforms:
        color = PLATFORM_COLORS.get(plat, "29405C")
        label = PLATFORM_LABELS.get(plat, plat)
        num_cols = len(months) + 1
        ws.merge_cells(
            start_row=header_row,
            start_column=col_offset,
            end_row=header_row,
            end_column=col_offset + num_cols - 1,
        )
        cell = ws.cell(row=header_row, column=col_offset, value=f"Citations in {label}")
        _style_header(cell, color)
        col_offset += num_cols

    sub_header_row = 5
    col_offset = 3
    for plat in platforms:
        color = PLATFORM_COLORS.get(plat, "29405C")
        for month in months:
            cell = ws.cell(row=sub_header_row, column=col_offset, value=MONTH_NAMES[month])
            _style_header(cell, color)
            col_offset += 1
        cell = ws.cell(row=sub_header_row, column=col_offset, value="Total")
        _style_header(cell, color, bold=True)
        col_offset += 1

    # Data rows
    data_start_row = 6
    for i, row in enumerate(rows):
        r = data_start_row + i
        ws.cell(row=r, column=1, value=row["sr_no"])
        _style_cell(ws.cell(row=r, column=1))

        ws.cell(row=r, column=2, value=row["faculty_name"])
        _style_cell(ws.cell(row=r, column=2), align="left")

        col_offset = 3
        for plat in platforms:
            plat_data = row.get(plat, {})
            for month in months:
                val = plat_data.get(month)
                cell = ws.cell(row=r, column=col_offset, value="NA" if val is None else val)
                _style_cell(ws.cell(row=r, column=col_offset))
                col_offset += 1
            total_val = plat_data.get("total", 0)
            cell = ws.cell(row=r, column=col_offset, value=total_val)
            _style_cell(ws.cell(row=r, column=col_offset), bold=True)
            col_offset += 1

    # Total row
    total_row = data_start_row + len(rows)
    ws.cell(row=total_row, column=2, value="Total")
    _style_cell(ws.cell(row=total_row, column=2), bold=True, align="left")

    col_offset = 3
    for plat in platforms:
        plat_totals = totals.get(plat, {})
        for month in months:
            cell = ws.cell(row=total_row, column=col_offset, value=plat_totals.get(month, 0))
            _style_cell(ws.cell(row=total_row, column=col_offset), bold=True)
            col_offset += 1
        cell = ws.cell(row=total_row, column=col_offset, value=plat_totals.get("total", 0))
        _style_cell(ws.cell(row=total_row, column=col_offset), bold=True)
        col_offset += 1

    # Average row
    avg_row = total_row + 1
    ws.cell(row=avg_row, column=2, value="Average")
    _style_cell(ws.cell(row=avg_row, column=2), bold=True, align="left")

    col_offset = 3
    for plat in platforms:
        plat_avgs = averages.get(plat, {})
        for month in months:
            cell = ws.cell(row=avg_row, column=col_offset, value=plat_avgs.get(month, 0))
            _style_cell(ws.cell(row=avg_row, column=col_offset), bold=True)
            col_offset += 1
        cell = ws.cell(row=avg_row, column=col_offset, value=plat_avgs.get("total", 0))
        _style_cell(ws.cell(row=avg_row, column=col_offset), bold=True)
        col_offset += 1

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    for c in range(3, col_offset):
        ws.column_dimensions[get_column_letter(c)].width = 9

    ws.freeze_panes = "C6"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_date_range_excel(report: dict) -> bytes:
    """Generate Excel file for custom date range report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Date Range Citations"

    platforms = report["platforms"]
    rows = report["rows"]
    totals = report["totals"]
    averages = report["averages"]
    dept = report["department"]

    from_dt = datetime.fromisoformat(report["from_date"])
    to_dt = datetime.fromisoformat(report["to_date"])

    # Title
    total_cols = 2 + len(platforms) * 2  # Sr, Name, then citations+papers per platform
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_cell = ws.cell(row=1, column=1, value=f"{dept} Department")
    title_cell.font = Font(bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    subtitle = ws.cell(row=2, column=1, value=f"Citation Report — {from_dt.strftime('%d %B %Y')} to {to_dt.strftime('%d %B %Y')}")
    subtitle.font = Font(bold=True, size=12, color="555555")
    subtitle.alignment = Alignment(horizontal="center")

    # Headers
    header_row = 4
    ws.cell(row=header_row, column=1, value="Sr. No.")
    _style_header(ws.cell(row=header_row, column=1), "29405C")
    ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row + 1, end_column=1)

    ws.cell(row=header_row, column=2, value="Faculty Name")
    _style_header(ws.cell(row=header_row, column=2), "29405C")
    ws.merge_cells(start_row=header_row, start_column=2, end_row=header_row + 1, end_column=2)

    col_offset = 3
    for plat in platforms:
        color = PLATFORM_COLORS.get(plat, "29405C")
        label = PLATFORM_LABELS.get(plat, plat)
        ws.merge_cells(start_row=header_row, start_column=col_offset, end_row=header_row, end_column=col_offset + 1)
        cell = ws.cell(row=header_row, column=col_offset, value=label)
        _style_header(cell, color)
        col_offset += 2

    sub_header_row = 5
    col_offset = 3
    for plat in platforms:
        color = PLATFORM_COLORS.get(plat, "29405C")
        cell = ws.cell(row=sub_header_row, column=col_offset, value="Papers")
        _style_header(cell, color)
        cell = ws.cell(row=sub_header_row, column=col_offset + 1, value="Citations")
        _style_header(cell, color)
        col_offset += 2

    # Data rows
    data_start_row = 6
    for i, row in enumerate(rows):
        r = data_start_row + i
        ws.cell(row=r, column=1, value=row["sr_no"])
        _style_cell(ws.cell(row=r, column=1))

        ws.cell(row=r, column=2, value=row["faculty_name"])
        _style_cell(ws.cell(row=r, column=2), align="left")

        col_offset = 3
        for plat in platforms:
            plat_data = row.get(plat, {})
            papers_val = plat_data.get("papers")
            cites_val = plat_data.get("citations")
            cell = ws.cell(row=r, column=col_offset, value="NA" if papers_val is None else papers_val)
            _style_cell(ws.cell(row=r, column=col_offset))
            cell = ws.cell(row=r, column=col_offset + 1, value="NA" if cites_val is None else cites_val)
            _style_cell(ws.cell(row=r, column=col_offset + 1))
            col_offset += 2

    # Total row
    total_row = data_start_row + len(rows)
    ws.cell(row=total_row, column=2, value="Total")
    _style_cell(ws.cell(row=total_row, column=2), bold=True, align="left")

    col_offset = 3
    for plat in platforms:
        plat_totals = totals.get(plat, {})
        cell = ws.cell(row=total_row, column=col_offset, value=plat_totals.get("papers", 0))
        _style_cell(ws.cell(row=total_row, column=col_offset), bold=True)
        cell = ws.cell(row=total_row, column=col_offset + 1, value=plat_totals.get("citations", 0))
        _style_cell(ws.cell(row=total_row, column=col_offset + 1), bold=True)
        col_offset += 2

    # Average row
    avg_row = total_row + 1
    ws.cell(row=avg_row, column=2, value="Average")
    _style_cell(ws.cell(row=avg_row, column=2), bold=True, align="left")

    col_offset = 3
    for plat in platforms:
        plat_avgs = averages.get(plat, {})
        cell = ws.cell(row=avg_row, column=col_offset, value=plat_avgs.get("papers", 0))
        _style_cell(ws.cell(row=avg_row, column=col_offset), bold=True)
        cell = ws.cell(row=avg_row, column=col_offset + 1, value=plat_avgs.get("citations", 0))
        _style_cell(ws.cell(row=avg_row, column=col_offset + 1), bold=True)
        col_offset += 2

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    for c in range(3, col_offset):
        ws.column_dimensions[get_column_letter(c)].width = 12

    ws.freeze_panes = "C6"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
